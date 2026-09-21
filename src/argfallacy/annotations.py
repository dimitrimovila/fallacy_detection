"""The annotations table, and the reader of the ``Annotazione Dumitru`` sheets.

``data/items.csv`` and ``data/annotations.csv`` were frozen once from the
workbook and from the prior runs (``docs/dati.md``).  Since then only the
``Annotazione Dumitru`` sheets move.  Each of them carries an ``item_id`` column,
so a row finds its item by the identifier and never by the text; the text serves
only as a check that the column was pasted on the right rows.

``argfallacy annotations update`` reads every sheet whose name starts with
``Annotazione Dumitru`` and replaces the ``dumitru`` rows of those sheets in the
table.
"""

from __future__ import annotations

import re
from difflib import SequenceMatcher
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .labels import (
    load_scheme_vocabulary,
    load_vocabulary,
    normalize_label,
    normalize_scheme_label,
)
from .labels.vocabulary import light_normalize
from .paths import REPO_ROOT
from .schemes import load_all
from .schemes.loader import SchemeError

ITEMS_FILE = REPO_ROOT / "data" / "items.csv"
ANNOTATIONS_FILE = REPO_ROOT / "data" / "annotations.csv"
COLUMNS = ["item_id", "annotator", "sheet", "row", "field", "value", "raw"]
SORT = ["item_id", "sheet", "row", "annotator", "field"]

SHEET_PREFIX = "Annotazione Dumitru"
ANNOTATOR = "dumitru"
SKIP = "ESCLUSO"
CQ_HEADER = re.compile(r"^CQ\d+(\.\d+)*$")
# `-` and `na`: the diagram did not reach the CQ; `n.a.(C=A)`: the conclusion is the
# assertion, so the CQ does not apply
CQ_ANSWERS = {"yes": "yes", "no": "no", "-": "na", "na": "na", "n.a.(c=a)": "na",
              "?": "idk", "idk": "idk", "positive": "positive", "negative": "negative"}


def load_annotations(path: str | Path = ANNOTATIONS_FILE) -> pd.DataFrame:
    """The table as strings.  ``na`` and ``none`` are values here, never missing cells."""
    frame = pd.read_csv(path, dtype=str, keep_default_na=False)
    frame["row"] = frame["row"].astype(int)
    return frame


def _same_text(a: str, b: str) -> bool:
    """Tolerates a copying slip, not a row pasted next to another item."""
    a, b = light_normalize(a), light_normalize(b)
    return a in b or b in a or SequenceMatcher(None, a, b, autojunk=False).ratio() >= 0.9


def read_sheet(sheet, texts: dict[str, str], schemes, vocab, scheme_vocab) -> list[dict]:
    """The ``dumitru`` rows of one sheet.  Anything the rules below cannot read stops here."""
    table = [["" if c is None else str(c).strip() for c in row]
             for row in sheet.iter_rows(values_only=True)]
    header, name = table[0], sheet.title

    def column(label: str, last: bool = False) -> int | None:
        found = [i for i, h in enumerate(header) if h == label]
        return (found[-1] if last else found[0]) if found else None

    cqs = {h: i for i, h in enumerate(header) if CQ_HEADER.match(h)}
    matching = [k for k, s in schemes.items() if set(s.cqs) == set(cqs)]
    if len(matching) != 1:
        raise SchemeError(f"{name}: the CQ columns {sorted(cqs)} are not those of one scheme")
    scheme = matching[0]
    where = {"item_id": column("item_id"), "text": column("text"),
             "scheme": column("scheme", last=True), "verdict": column("verdetto"),
             "uncertainty": column("incertezza"), "note": column("nota")}
    missing = [k for k in ("item_id", "text", "scheme", "verdict") if where[k] is None]
    if missing:
        raise SchemeError(f"{name}: no column for {missing}")

    idk = {light_normalize(m) for m in vocab["idk_markers"]}
    idk |= {light_normalize(m) for m in scheme_vocab["special_values"]["idk"]}
    absent = {light_normalize(m) for m in vocab["scheme_absent_markers"]}
    out, problems = [], []
    for number, row in enumerate(table[1:], start=2):
        def cell(i, row=row):
            return row[i] if i is not None and i < len(row) else ""

        ident, text = cell(where["item_id"]), cell(where["text"])
        if ident == SKIP or not (ident or text):
            continue
        if not ident:
            problems.append(f"row {number}: a text with no item_id")
            continue
        if ident not in texts:
            problems.append(f"row {number}: item_id {ident!r} is not in items.csv")
            continue
        if not text or not _same_text(text, texts[ident]):
            problems.append(f"row {number}: the text is not that of item {ident}")
            continue

        def add(field, value, raw, number=number, ident=ident):
            out.append({"item_id": ident, "annotator": ANNOTATOR, "sheet": name,
                        "row": number, "field": field, "value": value, "raw": raw})

        raw = cell(where["scheme"])
        key = light_normalize(raw)
        if key and CQ_ANSWERS.get(key) != "na":
            value = ("idk" if key in idk else scheme if key == "yes"
                     else "none" if key == "no" else normalize_scheme_label(raw, scheme_vocab))
            add("scheme", value, raw)
        raw = cell(where["verdict"])
        if light_normalize(raw) in absent:
            if light_normalize(cell(where["scheme"])) != "no":
                problems.append(f"row {number}: `{raw}` but the scheme is not NO")
        elif raw:
            add("verdict", normalize_label(raw, vocab).id, raw)
        for cq_id, i in cqs.items():
            if raw := cell(i):
                if light_normalize(raw) not in CQ_ANSWERS:
                    problems.append(f"row {number} {cq_id}: unknown answer {raw!r}")
                add(cq_id, CQ_ANSWERS.get(light_normalize(raw)), raw)
        for field in ("uncertainty", "note"):
            if raw := cell(where[field]):
                add(field, raw, raw)
    if problems:
        raise SchemeError(f"{name}:\n  - " + "\n  - ".join(problems))
    return out


def update(workbook: str | Path, path: str | Path = ANNOTATIONS_FILE,
           items: str | Path = ITEMS_FILE) -> dict[str, tuple[int, int]]:
    """Replace the ``dumitru`` rows of every sheet read.  Returns sheet -> (rows before, after)."""
    texts = dict(pd.read_csv(items, dtype=str, keep_default_na=False)[["item_id", "text"]].values)
    schemes, vocab = load_all(), load_vocabulary(check_schemes=False)
    scheme_vocab = load_scheme_vocabulary()
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        new = {s.title: read_sheet(s, texts, schemes, vocab, scheme_vocab)
               for s in book.worksheets if s.title.startswith(SHEET_PREFIX)}
    finally:
        book.close()
    table = load_annotations(path)
    mine = (table["annotator"] == ANNOTATOR) & table["sheet"].isin(new)
    before = table[mine]["sheet"].value_counts()
    rows = [r for sheet_rows in new.values() for r in sheet_rows]
    table = pd.concat([table[~mine], pd.DataFrame(rows, columns=COLUMNS)], ignore_index=True)
    table.sort_values(SORT).to_csv(path, index=False, lineterminator="\n")
    return {s: (int(before.get(s, 0)), len(r)) for s, r in new.items()}


__all__ = ["ANNOTATIONS_FILE", "ITEMS_FILE", "COLUMNS", "load_annotations",
           "read_sheet", "update"]
